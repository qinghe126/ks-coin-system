from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
import json
import os
import re
import time
import hmac
import hashlib
import base64
import urllib.parse
from datetime import datetime, timedelta
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import sqlite3
from werkzeug.security import generate_password_hash, check_password_hash
import threading
from collections import deque
import logging
import io

# 配置日志
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

app = Flask(__name__)
app.secret_key = 'ks-coin-panel-secret-key-2024-updated'
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=7)

# 全局配置
CONFIG = {
    'COIN_TO_MONEY_RATIO': 0.0001,  # 10000金币 = 1元
    'DEFAULT_COIN_THRESHOLD': 10000,
    'HISTORY_DAYS': 7
}

# 全局数据存储
query_history = {}
account_data = {}

# 数据库初始化状态
db_initialized = False

def init_db():
    """初始化数据库"""
    global db_initialized
    if db_initialized:
        return
        
    conn = sqlite3.connect('ks_data.db')
    c = conn.cursor()
    
    # 用户表
    c.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    
    # 青龙配置表
    c.execute('''
        CREATE TABLE IF NOT EXISTS qinglong_config (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            ql_url TEXT,
            client_id TEXT,
            client_secret TEXT,
            remark TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')
    
    # 账号配置表
    c.execute('''
        CREATE TABLE IF NOT EXISTS accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            remark TEXT,
            cookie TEXT NOT NULL,
            salt TEXT,
            user_agent TEXT,
            proxy_info TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')
    
    # 查询历史表
    c.execute('''
        CREATE TABLE IF NOT EXISTS query_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            account_id INTEGER,
            query_time TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            coin_balance INTEGER,
            cash_balance REAL,
            today_coin INTEGER,
            yesterday_coin INTEGER,
            total_balance REAL,
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (account_id) REFERENCES accounts (id)
        )
    ''')
    
    # 变量名称历史表
    c.execute('''
        CREATE TABLE IF NOT EXISTS var_name_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            var_name TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')
    
    # 插入默认管理员用户
    try:
        c.execute('INSERT INTO users (username, password_hash) VALUES (?, ?)',
                 ('admin', generate_password_hash('admin123')))
        logging.info("创建默认管理员账户: admin/admin123")
    except sqlite3.IntegrityError:
        logging.info("管理员账户已存在")
    
    conn.commit()
    conn.close()
    db_initialized = True
    logging.info("数据库初始化完成")

# 使用 before_request 替代 before_first_request
@app.before_request
def initialize_database():
    """在每个请求之前检查并初始化数据库"""
    if not db_initialized:
        init_db()

# 工具函数
def parse_cookie(txt):
    """解析Cookie字符串为字典"""
    try:
        if not txt:
            return {}
        if not txt.endswith(';'):
            txt += ';'
        pairs = txt.split(";")[:-1]
        ck_json = {}
        for pair in pairs:
            if '=' in pair:
                key, value = pair.split("=", 1)
                ck_json[key.strip()] = value.strip()
        return ck_json
    except Exception as e:
        logging.error(f"解析Cookie失败: {e}")
        return {}

def get_today_date():
    """获取当天日期"""
    today = datetime.now()
    return today.strftime("%Y.%-m.%-d")

def parse_accounts_from_text(text):
    """从文本解析账号配置"""
    accounts = []
    if not text:
        return accounts

    # 支持多种分隔符
    if '&' in text:
        account_list = text.split('&')
    elif '\n' in text:
        account_list = text.split('\n')
    else:
        account_list = [text]

    for i, acc in enumerate(account_list):
        acc = acc.strip()
        if not acc:
            continue

        account = {
            "id": i + 1,
            "name": f"账号{i+1}",
            "cookie": "",
            "salt": None,
            "ua": None,
            "proxy": None
        }

        # 检测格式类型
        if "@" in acc and acc.count("@") >= 3:
            parts = [p.strip() for p in acc.split("@")]
            account["name"] = parts[0] if parts[0] else account["name"]
            account["cookie"] = parts[1] if len(parts) > 1 and parts[1] else ""
            account["salt"] = parts[2] if len(parts) > 2 and parts[2] else None
            account["ua"] = parts[3] if len(parts) > 3 and parts[3] else None
            
        elif "#" in acc:
            parts = [p.strip() for p in acc.split("#")]
            account["name"] = parts[0] if parts[0] else account["name"]
            account["cookie"] = parts[1] if len(parts) > 1 and parts[1] else ""
            account["salt"] = parts[2] if len(parts) > 2 and parts[2] else None
            
        else:
            account["name"] = f"账号{i+1}"
            account["cookie"] = acc

        # 验证必要字段
        if not account["cookie"]:
            logging.warning(f"账号 {account['name']} 缺少cookie，跳过")
            continue

        accounts.append(account)

    logging.info(f"成功解析 {len(accounts)} 个账号")
    return accounts

# 青龙接口类 - 增强版
class QinglongAPI:
    def __init__(self, url, client_id, client_secret):
        self.base_url = url.rstrip('/')
        self.client_id = client_id
        self.client_secret = client_secret
        self.token = None
        self.token_expire = None
    
    def get_token(self):
        """获取青龙面板token"""
        if self.token and self.token_expire and datetime.now() < self.token_expire:
            return self.token
            
        url = f"{self.base_url}/open/auth/token"
        params = {
            "client_id": self.client_id,
            "client_secret": self.client_secret
        }
        
        try:
            response = requests.get(url, params=params, timeout=10)
            if response.status_code == 200:
                data = response.json()
                if data.get('code') == 200:
                    self.token = data['data']['token']
                    self.token_expire = datetime.now() + timedelta(minutes=110)
                    return self.token
            logging.error(f"获取青龙token失败: {response.text}")
            return None
        except Exception as e:
            logging.error(f"获取青龙token异常: {e}")
            return None
    
    def get_environments(self):
        """获取环境变量"""
        token = self.get_token()
        if not token:
            return None
            
        url = f"{self.base_url}/open/envs"
        headers = {
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json"
        }
        
        try:
            response = requests.get(url, headers=headers, timeout=10)
            if response.status_code == 200:
                data = response.json()
                if data.get('code') == 200:
                    return data['data']
            logging.error(f"获取环境变量失败: {response.text}")
            return None
        except Exception as e:
            logging.error(f"获取环境变量异常: {e}")
            return None
    
    def get_ks_cookies_by_name(self, var_name=None):
        """根据变量名称获取快手Cookie环境变量"""
        envs = self.get_environments()
        if not envs:
            return []
        
        ks_cookies = []
        keywords = ['kuaishou', 'ks', '快手', 'ksjsb', 'Lindong_ksjsb']
        
        for env in envs:
            name = env.get('name', '') or ''
            value = env.get('value', '') or ''
            remarks = env.get('remarks', '') or ''
            
            # 如果指定了变量名称，则精确匹配
            if var_name and var_name.strip():
                if var_name.strip().lower() == name.lower() or var_name.strip().lower() == remarks.lower():
                    ks_cookies.append({
                        'name': name,
                        'value': value,
                        'remarks': remarks,
                        'id': env.get('id')
                    })
            else:
                # 否则使用关键词匹配
                name_str = str(name) if name is not None else ''
                remarks_str = str(remarks) if remarks is not None else ''
                
                for keyword in keywords:
                    if (keyword.lower() in name_str.lower() or 
                        keyword.lower() in remarks_str.lower()):
                        ks_cookies.append({
                            'name': name_str,
                            'value': value,
                            'remarks': remarks_str,
                            'id': env.get('id')
                        })
                        break
        
        logging.info(f"从青龙获取到 {len(ks_cookies)} 个快手Cookie")
        return ks_cookies

# 快手客户端类
class KuaishouClient:
    def __init__(self, account_info):
        self.account_info = account_info
        self.coin_to_money_ratio = CONFIG['COIN_TO_MONEY_RATIO']  # 10000金币 = 1元
        
        self.nickName = "未知"
        self.user_id = "未知"
        self.balance = 0
        self.coin_balance = 0
        self.all_balance = 0
        self.today_coin = 0
        self.yesterday_coin = 0
        
        self.host = "nebula.kuaishou.com"
        self.base_url = f"https://{self.host}"
        self.headers = {
            'Host': self.host,
            'Connection': 'keep-alive',
            'User-Agent': self.account_info.get("ua") or 'Mozilla/5.0 (Linux; Android 13; PJA110 Build/TP1A.220905.001; wv) AppleWebKit/537.36 (KHTML, like Gecko) Version/4.0 Chrome/103.0.5060.129 Mobile Safari/537.36 Yoda/3.2.9-rc9 ksNebula/13.1.40.9558 OS_PRO_BIT/64 MAX_PHY_MEM/15190 KDT/PHONE AZPREFIX/az2 ICFO/0 StatusHT/31 TitleHT/44 NetType/WIFI ISLP/0 ISDM/0 ISLB/0 locale/zh-cn SHP/2606 SWP/1240 SD/3.5 CT/0 ISLM/0',
            'content-type': 'application/x-www-form-urlencoded;charset=UTF-8',
            'Accept': '*/*',
            'X-Requested-With': 'com.kuaishou.nebula',
            'Sec-Fetch-Site': 'same-origin',
            'Sec-Fetch-Mode': 'cors',
            'Sec-Fetch-Dest': 'empty',
            'Accept-Encoding': 'gzip, deflate',
            'Accept-Language': 'zh-CN,zh;q=0.9,en-US;q=0.8,en;q=0.7',
            'Cookie': self.account_info["cookie"]
        }
        
        # 解析用户ID
        try:
            user_id_match = re.search(r'userId\s*=\s*([^;]+)', self.account_info["cookie"], re.IGNORECASE)
            if user_id_match:
                self.user_id = user_id_match.group(1).strip()
        except Exception as e:
            logging.warning(f"解析用户ID失败: {e}")
            self.user_id = "未知"

    def send_request(self, api, headers=None, params=None, method="get"):
        """发送请求"""
        if headers is None:
            headers = self.headers
        if params is None:
            params = {}
            
        url = f"{self.base_url}{api}"
        try:
            if method == "get":
                response = requests.get(url, headers=headers, params=params, timeout=15, verify=False)
            else:
                response = requests.post(url, headers=headers, params=params, timeout=15, verify=False)
            return response.json()
        except Exception as e:
            logging.error(f"请求失败: {str(e)}")
            return None

    def _handle_response(self, res, code_key="result", code_value=1, data_key="data"):
        """处理响应"""
        if not res:
            return False, "接口请求失败或响应为空"
        res_code = res.get(code_key, "")
        if res_code != code_value:
            error_msg = res.get("msg", "未知错误")
            return False, error_msg
        result_data = res.get(data_key, {})
        return True, result_data

    def query_basic(self):
        """查询基本信息"""
        api = "/rest/n/nebula/activity/earn/overview/basicInfo"
        params = {"source": "timer"}
        res = self.send_request(api, params=params, method="get")
        success, result = self._handle_response(res)
        if success:
            self.balance = result.get('totalCash', 0)
            self.coin_balance = result.get('totalCoin', 0)
            user_info = result.get('userData', {})
            self.nickName = user_info.get('nickname', '未知')
            return True
        return False

    def query_detail(self):
        """查询详细信息"""
        api = "/rest/n/nebula/account/overview"
        params = {}
        res = self.send_request(api, params=params, method="get")
        success, result = self._handle_response(res)
        
        if success:
            try:
                self.coin_balance = int(result.get('coinBalance', 0))
                self.balance = float(result.get('cashBalance', 0))
                self.all_balance = float(result.get('accumulativeAmount', 0))
                
                # 获取今日和昨日日期
                today = get_today_date()
                yesterday = (datetime.now() - timedelta(days=1)).strftime("%Y.%-m.%-d")
                
                # 统计今日和昨日金币
                self.today_coin = 0
                self.yesterday_coin = 0
                
                recent_coin_records = []
                recent_cash_records = []
                
                # 处理金币记录
                coin_account = result.get('coinAccountPage', {})
                if coin_account:
                    for record in coin_account.get('data', []):
                        amount = int(record.get('amount', 0))
                        if amount <= 0:
                            continue
                        create_time = record.get('createTime', '')
                        if create_time == today:
                            self.today_coin += amount
                        elif create_time == yesterday:
                            self.yesterday_coin += amount
                        
                        if len(recent_coin_records) < 5:
                            recent_coin_records.append({
                                'time': create_time,
                                'type': record.get('eventType', '未知类型'),
                                'amount': amount,
                                'display': f"+{amount}金币"
                            })
                
                # 处理现金记录
                cash_account = result.get('cashAccountPage', {})
                if cash_account:
                    for record in cash_account.get('data', []):
                        if len(recent_cash_records) < 5:
                            try:
                                amount = float(record.get('amount', 0))
                                event_type = record.get('eventType', '未知类型')
                                sign = "+" if event_type != "金币兑换现金" and amount > 0 else ""
                                recent_cash_records.append({
                                    'time': record.get('createTime', ''),
                                    'type': event_type,
                                    'amount': amount,
                                    'display': f"{sign}{amount}元"
                                })
                            except:
                                continue
                
                return {
                    'success': True,
                    'nickname': self.nickName,
                    'user_id': self.user_id,
                    'remark': self.account_info['name'],
                    'coin_balance': self.coin_balance,
                    'cash_balance': self.balance,
                    'total_balance': self.all_balance,
                    'today_coin': self.today_coin,
                    'yesterday_coin': self.yesterday_coin,
                    'today_money': round(self.today_coin * self.coin_to_money_ratio, 2),
                    'yesterday_money': round(self.yesterday_coin * self.coin_to_money_ratio, 2),
                    'exchange_mode': "自动兑换" if result.get('accountState', 'NORMAL') == "NORMAL" else "手动兑换",
                    'recent_coin_records': recent_coin_records,
                    'recent_cash_records': recent_cash_records
                }
            except Exception as e:
                logging.error(f"处理查询结果异常: {e}")
                return {
                    'success': False,
                    'error': f'数据处理失败: {str(e)}'
                }
        else:
            return {
                'success': False,
                'error': result if isinstance(result, str) else '查询失败'
            }

# Excel生成函数 - 修改为返回文件流
def create_excel_summary(data_list):
    """创建Excel汇总表格并返回文件流"""
    coin_to_money_ratio = CONFIG['COIN_TO_MONEY_RATIO']
    
    wb = Workbook()
    ws = wb.active
    ws.title = "金币数据汇总"
    
    headers = ["序号", "用户名", "备注", "今日金币", "昨日金币", "总金币", 
               "今日金额(元)", "昨日金额(元)"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    
    total_today_coin = 0
    total_yesterday_coin = 0
    total_today_money = 0.0
    total_yesterday_money = 0.0
    
    for row, data in enumerate(data_list, 2):
        today_money = data["today_coin"] * coin_to_money_ratio
        yesterday_money = data["yesterday_coin"] * coin_to_money_ratio
        
        total_today_coin += data["today_coin"]
        total_yesterday_coin += data["yesterday_coin"]
        total_today_money += today_money
        total_yesterday_money += yesterday_money
        
        ws.cell(row=row, column=1, value=data["index"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=data["username"])
        ws.cell(row=row, column=3, value=data["remark"])
        ws.cell(row=row, column=4, value=data["today_coin"]).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=5, value=data["yesterday_coin"]).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=6, value=data["total_coin"]).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=7, value=round(today_money, 2)).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=8, value=round(yesterday_money, 2)).alignment = Alignment(horizontal="right")
    
    # 添加汇总行
    total_row = len(data_list) + 2
    ws.cell(row=total_row, column=1, value="累计总和").font = Font(bold=True)
    ws.merge_cells(f'A{total_row}:C{total_row}')
    ws.cell(row=total_row, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=total_row, column=4, value=total_today_coin).font = Font(bold=True)
    ws.cell(row=total_row, column=4).alignment = Alignment(horizontal="right")
    ws.cell(row=total_row, column=5, value=total_yesterday_coin).font = Font(bold=True)
    ws.cell(row=total_row, column=5).alignment = Alignment(horizontal="right")
    ws.cell(row=total_row, column=7, value=round(total_today_money, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=7).alignment = Alignment(horizontal="right")
    ws.cell(row=total_row, column=8, value=round(total_yesterday_money, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=8).alignment = Alignment(horizontal="right")
    
    # 调整列宽
    for col, width in zip('ABCDEFGH', [6, 20, 20, 12, 12, 12, 15, 15]):
        ws.column_dimensions[col].width = width
    
    # 将工作簿保存到内存中
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    return file_stream, total_today_coin, total_today_money, total_yesterday_coin, total_yesterday_money

# 路由定义
@app.route('/')
def index():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    return render_template('index.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'GET':
        return render_template('login.html')
    
    username = request.form.get('username')
    password = request.form.get('password')
    
    if not username or not password:
        return jsonify({'success': False, 'message': '用户名和密码不能为空'})
    
    conn = sqlite3.connect('ks_data.db')
    c = conn.cursor()
    c.execute('SELECT id, password_hash FROM users WHERE username = ?', (username,))
    user = c.fetchone()
    conn.close()
    
    if user and check_password_hash(user[1], password):
        session['user_id'] = user[0]
        session['username'] = username
        session.permanent = True
        logging.info(f"用户 {username} 登录成功")
        return jsonify({'success': True})
    else:
        logging.warning(f"用户 {username} 登录失败")
        return jsonify({'success': False, 'message': '用户名或密码错误'})

@app.route('/logout')
def logout():
    username = session.get('username', '未知用户')
    session.clear()
    logging.info(f"用户 {username} 退出登录")
    return redirect(url_for('login'))

@app.route('/query', methods=['POST'])
def query_accounts():
    if 'user_id' not in session:
        return jsonify({'success': False, 'message': '请先登录'})
    
    try:
        cookies_text = request.form.get('cookies', '')
        coin_threshold = int(request.form.get('coin_threshold', CONFIG['DEFAULT_COIN_THRESHOLD']))
        
        if not cookies_text:
            return jsonify({'success': False, 'message': '请输入Cookie数据'})
        
        # 解析账号
        accounts = parse_accounts_from_text(cookies_text)
        if not accounts:
            return jsonify({'success': False, 'message': '未解析到有效账号'})
        
        results = []
        summary_data = []
        
        for account in accounts:
            client = KuaishouClient(account)
            
            # 先查询基本信息获取昵称
            if client.query_basic():
                detail_result = client.query_detail()
                results.append(detail_result)
                
                if detail_result['success']:
                    summary_data.append({
                        "index": account["id"],
                        "username": detail_result['nickname'],
                        "remark": detail_result['remark'],
                        "today_coin": detail_result['today_coin'],
                        "yesterday_coin": detail_result['yesterday_coin'],
                        "total_coin": detail_result['coin_balance'],
                        "below_threshold": detail_result['today_coin'] < coin_threshold
                    })
            else:
                results.append({
                    'success': False,
                    'remark': account['name'],
                    'error': '基础查询失败'
                })
        
        # 计算汇总数据
        total_today_coin = sum(data["today_coin"] for data in summary_data)
        total_today_money = round(total_today_coin * CONFIG['COIN_TO_MONEY_RATIO'], 2)
        total_yesterday_coin = sum(data["yesterday_coin"] for data in summary_data)
        total_yesterday_money = round(total_yesterday_coin * CONFIG['COIN_TO_MONEY_RATIO'], 2)
        
        # 存储查询结果
        user_id = session['user_id']
        account_data[user_id] = {
            'results': results,
            'summary': summary_data,
            'query_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'total_today_coin': total_today_coin,
            'total_today_money': total_today_money,
            'total_yesterday_coin': total_yesterday_coin,
            'total_yesterday_money': total_yesterday_money
        }
        
        logging.info(f"用户 {session['username']} 查询了 {len(accounts)} 个账号，成功 {len(summary_data)} 个")
        
        return jsonify({
            'success': True,
            'results': results,
            'summary': {
                'total_accounts': len(accounts),
                'success_accounts': len([r for r in results if r.get('success')]),
                'total_today_coin': total_today_coin,
                'total_today_money': total_today_money,
                'total_yesterday_coin': total_yesterday_coin,
                'total_yesterday_money': total_yesterday_money,
                'coin_threshold': coin_threshold
            }
        })
        
    except Exception as e:
        logging.error(f"查询异常: {e}")
        return jsonify({'success': False, 'message': f'查询异常: {str(e)}'})

@app.route('/download_excel')
def download_excel():
    if 'user_id' not in session:
        return "请先登录"
    
    try:
        user_id = session['user_id']
        if user_id not in account_data:
            return "没有可下载的数据，请先执行查询"
        
        summary_data = account_data[user_id]['summary']
        if not summary_data:
            return "没有有效的数据可下载"
        
        # 生成Excel文件流
        file_stream, total_today_coin, total_today_money, total_yesterday_coin, total_yesterday_money = create_excel_summary(summary_data)
        
        filename = f"金币数据汇总_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return send_file(
            file_stream,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        
    except Exception as e:
        logging.error(f"下载Excel异常: {e}")
        return f"文件下载失败: {str(e)}"

@app.route('/get_history')
def get_history():
    if 'user_id' not in session:
        return jsonify({'success': False})
    
    user_id = session['user_id']
    if user_id in account_data:
        return jsonify({'success': True, 'data': account_data[user_id]})
    else:
        return jsonify({'success': True, 'data': None})

@app.route('/save_qinglong_config', methods=['POST'])
def save_qinglong_config():
    if 'user_id' not in session:
        return jsonify({'success': False, 'message': '请先登录'})
    
    try:
        ql_url = request.form.get('ql_url', '').strip()
        client_id = request.form.get('client_id', '').strip()
        client_secret = request.form.get('client_secret', '').strip()
        remark = request.form.get('remark', '').strip()
        
        if not ql_url or not client_id or not client_secret:
            return jsonify({'success': False, 'message': '请填写完整的青龙配置信息'})
        
        user_id = session['user_id']
        
        conn = sqlite3.connect('ks_data.db')
        c = conn.cursor()
        
        # 检查是否已存在配置
        c.execute('SELECT id FROM qinglong_config WHERE user_id = ?', (user_id,))
        existing = c.fetchone()
        
        if existing:
            c.execute('''
                UPDATE qinglong_config 
                SET ql_url = ?, client_id = ?, client_secret = ?, remark = ?
                WHERE user_id = ?
            ''', (ql_url, client_id, client_secret, remark, user_id))
            action = "更新"
        else:
            c.execute('''
                INSERT INTO qinglong_config (user_id, ql_url, client_id, client_secret, remark)
                VALUES (?, ?, ?, ?, ?)
            ''', (user_id, ql_url, client_id, client_secret, remark))
            action = "保存"
        
        conn.commit()
        conn.close()
        
        logging.info(f"用户 {session['username']} {action}青龙配置")
        return jsonify({'success': True, 'message': f'青龙配置{action}成功'})
        
    except Exception as e:
        logging.error(f"保存青龙配置异常: {e}")
        return jsonify({'success': False, 'message': f'保存配置失败: {str(e)}'})

@app.route('/get_qinglong_config')
def get_qinglong_config():
    if 'user_id' not in session:
        return jsonify({'success': False})
    
    user_id = session['user_id']
    
    conn = sqlite3.connect('ks_data.db')
    c = conn.cursor()
    c.execute('SELECT ql_url, client_id, client_secret, remark FROM qinglong_config WHERE user_id = ?', (user_id,))
    config = c.fetchone()
    conn.close()
    
    if config:
        return jsonify({
            'success': True,
            'config': {
                'ql_url': config[0],
                'client_id': config[1],
                'client_secret': config[2],
                'remark': config[3]
            }
        })
    else:
        return jsonify({'success': True, 'config': None})

@app.route('/get_qinglong_cookies_by_name', methods=['POST'])
def get_qinglong_cookies_by_name():
    """根据变量名称获取青龙Cookie"""
    if 'user_id' not in session:
        return jsonify({'success': False, 'message': '请先登录'})
    
    try:
        var_name = request.form.get('var_name', '').strip()
        
        user_id = session['user_id']
        
        conn = sqlite3.connect('ks_data.db')
        c = conn.cursor()
        c.execute('SELECT ql_url, client_id, client_secret FROM qinglong_config WHERE user_id = ?', (user_id,))
        config = c.fetchone()
        conn.close()
        
        if not config:
            return jsonify({'success': False, 'message': '请先配置青龙信息'})
        
        ql_url, client_id, client_secret = config
        
        # 根据变量名称获取青龙Cookie
        ql_api = QinglongAPI(ql_url, client_id, client_secret)
        cookies = ql_api.get_ks_cookies_by_name(var_name)
        
        if cookies is None:
            return jsonify({'success': False, 'message': '获取青龙Cookie失败，请检查配置'})
        
        # 格式化Cookie文本 - 只导入环境变量的值，并用&分隔
        cookie_text = ""
        for cookie in cookies:
            value = cookie.get('value', '').strip()
            if value:
                if cookie_text:  # 如果不是第一个账号，添加&分隔符
                    cookie_text += "&"
                cookie_text += value
        
        return jsonify({
            'success': True,
            'cookies': cookie_text,
            'count': len(cookies),
            'message': f'成功获取 {len(cookies)} 个账号' if len(cookies) > 0 else '未找到匹配的变量'
        })
        
    except Exception as e:
        logging.error(f"按名称获取青龙Cookie异常: {e}")
        return jsonify({'success': False, 'message': f'获取Cookie失败: {str(e)}'})

@app.route('/save_var_name', methods=['POST'])
def save_var_name():
    """保存变量名称到历史记录"""
    if 'user_id' not in session:
        return jsonify({'success': False, 'message': '请先登录'})
    
    try:
        var_name = request.form.get('var_name', '').strip()
        
        if not var_name:
            return jsonify({'success': False, 'message': '变量名称不能为空'})
        
        user_id = session['user_id']
        
        conn = sqlite3.connect('ks_data.db')
        c = conn.cursor()
        
        # 检查是否已存在
        c.execute('SELECT id FROM var_name_history WHERE user_id = ? AND var_name = ?', (user_id, var_name))
        existing = c.fetchone()
        
        if not existing:
            c.execute('INSERT INTO var_name_history (user_id, var_name) VALUES (?, ?)', (user_id, var_name))
        
        conn.commit()
        conn.close()
        
        return jsonify({'success': True, 'message': '变量名称保存成功'})
        
    except Exception as e:
        logging.error(f"保存变量名称异常: {e}")
        return jsonify({'success': False, 'message': f'保存失败: {str(e)}'})

@app.route('/get_var_names')
def get_var_names():
    """获取保存的变量名称列表"""
    if 'user_id' not in session:
        return jsonify({'success': False})
    
    user_id = session['user_id']
    
    conn = sqlite3.connect('ks_data.db')
    c = conn.cursor()
    c.execute('SELECT var_name FROM var_name_history WHERE user_id = ? ORDER BY created_at DESC', (user_id,))
    var_names = [row[0] for row in c.fetchall()]
    conn.close()
    
    return jsonify({'success': True, 'var_names': var_names})

@app.route('/delete_var_name', methods=['POST'])
def delete_var_name():
    """删除保存的变量名称"""
    if 'user_id' not in session:
        return jsonify({'success': False, 'message': '请先登录'})
    
    try:
        var_name = request.form.get('var_name', '').strip()
        
        if not var_name:
            return jsonify({'success': False, 'message': '变量名称不能为空'})
        
        user_id = session['user_id']
        
        conn = sqlite3.connect('ks_data.db')
        c = conn.cursor()
        c.execute('DELETE FROM var_name_history WHERE user_id = ? AND var_name = ?', (user_id, var_name))
        conn.commit()
        conn.close()
        
        return jsonify({'success': True, 'message': '变量名称删除成功'})
        
    except Exception as e:
        logging.error(f"删除变量名称异常: {e}")
        return jsonify({'success': False, 'message': f'删除失败: {str(e)}'})

@app.route('/get_config')
def get_config():
    return jsonify({
        'success': True,
        'coin_to_money_ratio': CONFIG['COIN_TO_MONEY_RATIO'],
        'default_threshold': CONFIG['DEFAULT_COIN_THRESHOLD']
    })

# 健康检查接口
@app.route('/health')
def health_check():
    return jsonify({
        'status': 'healthy',
        'timestamp': datetime.now().isoformat(),
        'version': '2.2.0'
    })

if __name__ == '__main__':
    # 禁用SSL警告
    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    
    # 初始化数据库
    init_db()
    
    # 启动应用
    app.run(host='0.0.0.0', port=5000, debug=False)